"""Report generation and export (CSV, Excel, PDF)."""
import csv
import io
from datetime import datetime

REPORT_TYPES = {
    "sales": "Sales register",
    "product_sales": "Product-wise sales & profit",
    "inventory": "Inventory status",
    "customers": "Customer summary",
}


def build_report(conn, kind, start, end):
    """Return {title, columns, rows, summary, money_cols} for a report."""
    if kind == "sales":
        data = conn.execute(
            """SELECT s.invoice_no, s.sale_date, COALESCE(c.name, 'Walk-in'),
                      (SELECT SUM(quantity) FROM sale_items WHERE sale_id = s.id),
                      s.payment_method, s.total
               FROM sales s LEFT JOIN customers c ON c.id = s.customer_id
               WHERE date(s.sale_date) BETWEEN ? AND ?
               ORDER BY s.sale_date DESC""", (start, end)).fetchall()
        columns = ["Invoice", "Date & time", "Customer", "Units", "Payment", "Total (INR)"]
        money = {5}
        revenue = sum(r[5] for r in data)
        summary = [("Orders", f"{len(data):,}"), ("Revenue", revenue),
                   ("Average order", revenue / len(data) if data else 0)]

    elif kind == "product_sales":
        data = conn.execute(
            """SELECT p.sku, p.name, p.category, SUM(si.quantity), SUM(si.line_total),
                      SUM(si.line_total - si.quantity * p.cost)
               FROM sale_items si
               JOIN sales s ON s.id = si.sale_id
               JOIN products p ON p.id = si.product_id
               WHERE date(s.sale_date) BETWEEN ? AND ?
               GROUP BY p.id ORDER BY SUM(si.line_total) DESC""", (start, end)).fetchall()
        columns = ["SKU", "Product", "Category", "Units sold", "Revenue (INR)", "Gross profit (INR)"]
        money = {4, 5}
        revenue, profit = sum(r[4] for r in data), sum(r[5] for r in data)
        summary = [("Revenue", revenue), ("Gross profit", profit),
                   ("Margin", f"{(profit / revenue * 100) if revenue else 0:.1f}%")]

    elif kind == "inventory":
        data = conn.execute(
            """SELECT sku, name, category, stock, reorder_level, price, stock * cost,
                      CASE WHEN stock = 0 THEN 'Out of stock'
                           WHEN stock <= reorder_level THEN 'Low' ELSE 'OK' END
               FROM products ORDER BY category, name""").fetchall()
        columns = ["SKU", "Product", "Category", "Stock", "Reorder level",
                   "Price (INR)", "Stock value (INR)", "Status"]
        money = {5, 6}
        summary = [("Products", f"{len(data):,}"),
                   ("Stock value (at cost)", sum(r[6] for r in data)),
                   ("Low / out of stock", f"{sum(1 for r in data if r[7] != 'OK')}")]

    elif kind == "customers":
        data = conn.execute(
            """SELECT c.name, c.phone, c.city, COUNT(s.id), COALESCE(SUM(s.total), 0),
                      COALESCE(MAX(date(s.sale_date)), '-')
               FROM customers c
               LEFT JOIN sales s ON s.customer_id = c.id
                    AND date(s.sale_date) BETWEEN ? AND ?
               GROUP BY c.id ORDER BY COALESCE(SUM(s.total), 0) DESC""", (start, end)).fetchall()
        columns = ["Customer", "Phone", "City", "Orders", "Total spent (INR)", "Last purchase"]
        money = {4}
        active = sum(1 for r in data if r[3])
        summary = [("Customers", f"{len(data):,}"), ("Active in period", f"{active:,}"),
                   ("Revenue from registered", sum(r[4] for r in data))]
    else:
        raise ValueError(f"Unknown report type: {kind}")

    return {
        "kind": kind,
        "title": REPORT_TYPES[kind],
        "columns": columns,
        "rows": [list(r) for r in data],
        "summary": summary,
        "money_cols": money,
        "start": start,
        "end": end,
    }


def _fmt(value, is_money):
    if is_money and isinstance(value, (int, float)):
        return f"{value:,.2f}"
    return "" if value is None else value


def to_csv(report):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(report["columns"])
    for r in report["rows"]:
        w.writerow([round(v, 2) if isinstance(v, float) else v for v in r])
    return io.BytesIO(buf.getvalue().encode("utf-8-sig"))


def to_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = report["kind"][:31]
    ws.append([report["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Period: {report['start']} to {report['end']}"])
    ws.append([])
    ws.append(report["columns"])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="14213D")
        cell.alignment = Alignment(horizontal="center")
    for r in report["rows"]:
        ws.append(r)
    for i in report["money_cols"]:
        for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=i + 1, max_col=i + 1):
            cell.number_format = "#,##0.00"
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col[header_row - 1:])
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 45)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def to_pdf(report, business_name="SmartBiz AI"):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=12 * mm,
                            rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"{business_name} - {report['title']}", styles["Title"]),
        Paragraph(f"Period: {report['start']} to {report['end']} &nbsp;&nbsp; "
                  f"Generated: {datetime.now():%d %b %Y %H:%M}", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    summary = " &nbsp;|&nbsp; ".join(
        f"<b>{k}:</b> {('Rs. ' + format(v, ',.2f')) if isinstance(v, (int, float)) else v}"
        for k, v in report["summary"])
    story += [Paragraph(summary, styles["Normal"]), Spacer(1, 4 * mm)]

    body = [report["columns"]] + [
        [_fmt(v, i in report["money_cols"]) for i, v in enumerate(r)] for r in report["rows"]
    ]
    table = Table(body, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14213D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F1EA")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9D5C9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(table)
    doc.build(story)
    out.seek(0)
    return out
